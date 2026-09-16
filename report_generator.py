from io import BytesIO
from pathlib import Path
from math import ceil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from PIL import Image

# Volaris-inspired palette. This is a visual approximation, not an official template.
MAGENTA = RGBColor(205, 0, 112)
DARK_MAGENTA = RGBColor(112, 0, 75)
LIGHT_MAGENTA = RGBColor(244, 220, 236)
DARK = RGBColor(45, 45, 45)
MID = RGBColor(105, 105, 105)
WHITE = RGBColor(255, 255, 255)
LIGHT_GRAY = RGBColor(242, 242, 242)
BORDER = RGBColor(210, 210, 210)

SECTION_CONFIG = {
    "data_plates": {"title": "ENG DATA PLATES", "layout": "standard"},
    "fwd_after": {"title": "FWD AND AFTER", "layout": "standard"},
    "cone": {"title": "FWD AND AFTER (CONE)", "layout": "standard"},
    "fan_blades": {"title": "FWD AND AFTER (FAN BLADES)", "layout": "standard"},
    "rubstrip": {"title": "FWD AND AFTER (RUBSTRIP PANEL, ACOUSTIC SEGMENT)", "layout": "standard"},
    "fan_section": {"title": "FAN SECTION", "layout": "standard"},
    "core": {"title": "CORE", "layout": "standard"},
    "upper_lower": {"title": "UPPER AND LOWER", "layout": "standard"},
    "fwd_mount": {"title": "UPPER AND LOWER (FWD MOUNT)", "layout": "standard"},
    "eec": {"title": "EEC DATA PLATE & SOFTWARE PLATE", "layout": "standard"},
    "dsu": {"title": "DSU", "layout": "standard"},
    "phmu": {"title": "PHMU", "layout": "standard"},
    "oil_tank": {"title": "OIL TANK", "layout": "standard"},
    "tarp": {"title": "PLASTIC TARP INSIDE-LOWER AREA AND DESICCANTS", "layout": "standard"},
    "stand": {"title": "Engine Stand & Cradle Data Plates", "layout": "standard"},
    "safety_pins": {"title": "SAFETY PINS ENG STAND", "layout": "standard"},
    "humidity": {"title": "HUMIDITY INDICATOR", "layout": "standard"},
    "covered": {"title": "ENG COVERED", "layout": "standard"},
    "tm": {"title": "TM VOI-TIC-111", "layout": "standard"},
}

SLIDE_W = 13.333
SLIDE_H = 7.5
OUT = Path("output")
OUT.mkdir(exist_ok=True)

def _text(slide, text, x, y, w, h, size=12, bold=False, color=DARK,
          align=PP_ALIGN.LEFT, font="Arial", valign=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box

def _rect(slide, x, y, w, h, fill=None, line=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill or WHITE
    shp.line.color.rgb = line or fill or WHITE
    return shp

def _add_header(slide, sn, section=None):
    # Thin magenta top rule + technical header inspired by the supplied report.
    _rect(slide, 0, 0, SLIDE_W, 0.08, MAGENTA)
    _text(
        slide,
        f"ENG INCOMING REPORT  S/N: {sn}",
        0.38, 0.15, 8.5, 0.38,
        size=12, bold=True, color=DARK
    )
    if section:
        _text(slide, section, 0.38, 0.58, 12.55, 0.38,
              size=16, bold=True, color=DARK)
    _rect(slide, 0.38, 0.98, 12.55, 0.025, MAGENTA)

def _add_footer(slide, page_no):
    _rect(slide, 0.38, 7.17, 12.55, 0.012, BORDER)
    _text(slide, f"{page_no:02d}", 12.2, 7.19, 0.7, 0.2,
          size=8, color=MID, align=PP_ALIGN.RIGHT)

def _image_bytes(obj):
    if hasattr(obj, "getvalue"):
        return obj.getvalue()
    return Path(obj).read_bytes()

def _add_image_contain(slide, image_obj, x, y, w, h, border=True):
    data = _image_bytes(image_obj)
    try:
        with Image.open(BytesIO(data)) as im:
            iw, ih = im.size
    except Exception:
        return

    ratio = iw / ih
    box_ratio = w / h
    if ratio >= box_ratio:
        rw = w
        rh = w / ratio
        rx = x
        ry = y + (h - rh) / 2
    else:
        rh = h
        rw = h * ratio
        rx = x + (w - rw) / 2
        ry = y

    if border:
        _rect(slide, x, y, w, h, WHITE, BORDER)
    slide.shapes.add_picture(
        BytesIO(data), Inches(rx), Inches(ry),
        width=Inches(rw), height=Inches(rh)
    )

def _add_cover(prs, pn, sn, qc, report_date):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    # Large editorial magenta block.
    _rect(slide, 0, 0, 4.1, SLIDE_H, MAGENTA)
    _text(slide, "ENG", 0.55, 1.05, 2.8, 0.75, 38, True, WHITE)
    _text(slide, "INCOMING", 0.55, 1.78, 3.0, 0.65, 25, True, WHITE)
    _text(slide, "ENGINE REPORT", 0.55, 2.42, 3.0, 0.45, 13, True, WHITE)

    _text(slide, "PART NUMBER", 4.75, 1.12, 2.1, 0.3, 9, True, MAGENTA)
    _text(slide, pn, 4.75, 1.43, 7.3, 0.55, 22, True, DARK)
    _text(slide, "SERIAL NUMBER", 4.75, 2.20, 2.1, 0.3, 9, True, MAGENTA)
    _text(slide, sn, 4.75, 2.51, 7.3, 0.55, 22, True, DARK)
    _text(slide, "QC", 4.75, 3.30, 1.0, 0.3, 9, True, MAGENTA)
    _text(slide, qc, 4.75, 3.60, 2.5, 0.45, 16, True, DARK)
    _text(slide, "DATE", 8.0, 3.30, 1.0, 0.3, 9, True, MAGENTA)
    _text(slide, report_date, 8.0, 3.60, 3.4, 0.45, 16, True, DARK)
    _rect(slide, 4.75, 4.45, 7.8, 0.05, MAGENTA)
    _text(
        slide,
        "ENGINE INCOMING INSPECTION",
        4.75, 4.75, 7.5, 0.45,
        size=15, bold=True, color=DARK
    )
    _text(
        slide,
        "Technical photographic report",
        4.75, 5.25, 7.5, 0.35,
        size=10, color=MID
    )
    _text(slide, "V2", 11.85, 6.78, 0.55, 0.25, 8, True, MAGENTA,
          align=PP_ALIGN.RIGHT)
    return slide

def _add_standard_section(prs, sn, title, photos, start_page):
    if not photos:
        return start_page

    per_page = 4
    total_pages = ceil(len(photos) / per_page)

    for page_index in range(total_pages):
        batch = photos[page_index * per_page:(page_index + 1) * per_page]
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _add_header(slide, sn, title)

        # 2x2 technical photo grid with generous white space.
        positions = [
            (0.55, 1.25, 6.0, 2.65),
            (6.78, 1.25, 6.0, 2.65),
            (0.55, 4.08, 6.0, 2.65),
            (6.78, 4.08, 6.0, 2.65),
        ]
        for photo, pos in zip(batch, positions):
            _add_image_contain(slide, photo, *pos)

        _add_footer(slide, start_page + page_index)
    return start_page + total_pages

def _add_discrepancy(prs, sn, item, page_no):
    photos = item.get("photos", [])
    description = item.get("description", "") or "DISCREPANCY"
    finding = item.get("finding", "") or ""
    reference = item.get("reference", "") or ""

    # If no photos, still create a documentation slide.
    if not photos:
        photos = []

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(slide, sn, "DISCREPANCIES")

    _text(slide, description, 0.55, 1.08, 12.0, 0.48, 14, True, MAGENTA)

    if photos:
        _add_image_contain(slide, photos[0], 0.55, 1.68, 7.25, 4.65)
        if len(photos) > 1:
            _add_image_contain(slide, photos[1], 8.05, 1.68, 4.75, 2.15)
        if len(photos) > 2:
            _add_image_contain(slide, photos[2], 8.05, 4.18, 4.75, 2.15)

    # Technical finding panel.
    panel_y = 6.42
    _rect(slide, 0.55, panel_y, 12.25, 0.55, LIGHT_MAGENTA, LIGHT_MAGENTA)
    _text(slide, "FINDING / OBSERVATION", 0.75, panel_y + 0.05, 2.1, 0.18,
          7.5, True, DARK_MAGENTA)
    _text(slide, finding, 2.7, panel_y + 0.03, 7.3, 0.23,
          8.5, False, DARK)
    if reference:
        _text(slide, reference, 10.0, panel_y + 0.03, 2.55, 0.28,
              8, True, DARK_MAGENTA, PP_ALIGN.RIGHT)

    _add_footer(slide, page_no)

def generate_report(pn, sn, qc, report_date, uploads, discrepancies=None):
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    _add_cover(prs, pn, sn, qc, report_date)

    page_no = 2
    for key, config in SECTION_CONFIG.items():
        page_no = _add_standard_section(
            prs, sn, config["title"], uploads.get(key, []), page_no
        )

    for item in discrepancies or []:
        if any([
            item.get("description"),
            item.get("finding"),
            item.get("reference"),
            item.get("photos"),
        ]):
            _add_discrepancy(prs, sn, item, page_no)
            page_no += 1

    # Final report-completed slide, echoing the supplied report's closing information.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(slide, sn, "REPORT COMPLETED")
    _text(slide, "REPORT ACCOMPLISHED BY QC", 0.7, 2.25, 5.0, 0.4,
          14, True, MAGENTA)
    _text(slide, qc, 0.7, 2.75, 5.0, 0.5, 20, True, DARK)
    _text(slide, "DATE", 0.7, 3.65, 2.0, 0.35, 10, True, MAGENTA)
    _text(slide, report_date, 0.7, 4.05, 4.5, 0.5, 18, True, DARK)
    _rect(slide, 0.7, 4.85, 5.2, 0.05, MAGENTA)
    _text(slide, "ENGINE INCOMING REPORT", 0.7, 5.15, 5.8, 0.4,
          12, True, DARK)
    _add_footer(slide, page_no)

    output = OUT / f"INCOMING_{sn}.pptx"
    prs.save(output)
    return output
